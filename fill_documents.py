"""Insert the computed numbers, sentences and figures into the manuscript and the response letter.

    python fill_documents.py --values results/main/analysis/values.json \
        --docs ../manuscript/information-4236555-revised.docx ../manuscript/response_to_reviewers.docx

Placeholders look like  ⟦KEY⟧  (numbers / sentences) and  ⟦FIG:name⟧  (a paragraph that becomes a figure).
Each document is written next to the original as *_FILLED.docx. Placeholders without a value are
listed at the end and left in place, so nothing is silently dropped.
"""
import argparse
import json
import os
import re

from docx import Document
from docx.shared import Cm

def _fix_settings(doc):
    from docx.oxml.ns import qn as _qn
    z = doc.settings.element.find(_qn("w:zoom"))
    if z is not None and z.get(_qn("w:percent")) is None:
        z.set(_qn("w:percent"), "100")


TOKEN = re.compile(r"⟦([A-Za-z0-9_:.\-]+)⟧")
FIG_WIDTH_CM = {"synthetic_examples": 12.0, "self_training": 14.0}


def iter_paragraphs(doc):
    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from walk(part)


def replace_in_paragraph(p, values, missing):
    runs = p.runs
    full = "".join(r.text for r in runs)
    if "⟦" not in full:
        return 0
    n = 0
    for m in reversed(list(TOKEN.finditer(full))):
        key = m.group(1)
        if key.startswith("FIG:"):
            continue
        if key not in values:
            missing.add(key)
            continue
        # locate the runs covering [m.start(), m.end())
        pos, first = 0, None
        for i, r in enumerate(runs):
            a, b = pos, pos + len(r.text)
            if first is None and a <= m.start() < b:
                first = i
                r.text = r.text[:m.start() - a] + str(values[key]) + (r.text[m.end() - a:] if m.end() <= b else "")
                if m.end() <= b:
                    break
            elif first is not None:
                if m.end() <= b:
                    r.text = r.text[m.end() - a:]
                    break
                r.text = ""
            pos = b
        n += 1
        full = "".join(r.text for r in runs)
    return n


def insert_figures(doc, fig_dir, missing):
    n = 0
    for p in iter_paragraphs(doc):
        m = TOKEN.fullmatch(p.text.strip())
        if not m or not m.group(1).startswith("FIG:"):
            continue
        name = m.group(1)[4:]
        path = os.path.join(fig_dir, f"fig_{name}.png")
        if not os.path.exists(path):
            missing.add(m.group(1))
            continue
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        p.runs[0].text = ""
        p.runs[0].add_picture(path, width=Cm(FIG_WIDTH_CM.get(name, 15.0)))
        n += 1
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--values", required=True)
    ap.add_argument("--docs", nargs="+", required=True)
    ap.add_argument("--figures", default=None, help="folder with fig_*.png (default: next to values.json)")
    ap.add_argument("--set", nargs="*", default=[], metavar="KEY=VALUE",
                    help="extra values, e.g. --set CODE_URL=https://github.com/<user>/<repo>")
    a = ap.parse_args()
    values = json.load(open(a.values))
    for kv in a.set:
        k, v = kv.split("=", 1)
        values[k] = v
    fig_dir = a.figures or os.path.dirname(os.path.abspath(a.values))
    for path in a.docs:
        doc = Document(path)
        missing = set()
        nf = insert_figures(doc, fig_dir, missing)
        nv = sum(replace_in_paragraph(p, values, missing) for p in iter_paragraphs(doc))
        out = os.path.splitext(path)[0] + "_FILLED.docx"
        _fix_settings(doc)
        doc.save(out)
        print(f"{os.path.basename(path)}: {nv} values and {nf} figures inserted -> {out}")
        if missing:
            print(f"  NOT FILLED ({len(missing)}): {', '.join(sorted(missing))}")


if __name__ == "__main__":
    main()
